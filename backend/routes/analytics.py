from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from database import get_db
from models import User, Video, ViewRecord
from auth import get_current_user
from datetime import datetime, timedelta
import io
import csv
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.platypus.frames import Frame
from reportlab.platypus.doctemplate import PageTemplate, BaseDocTemplate

router = APIRouter(prefix="/api/analytics", tags=["Analytics"])


def _safe_date(date_str: str):
    try:
        return datetime.strptime(date_str, "%Y-%m-%d").date()
    except Exception:
        return None


def _platform_label(url: str) -> str:
    url_l = (url or "").lower()
    if "instagram.com" in url_l:
        return "Instagram"
    if "facebook.com" in url_l or "fb.watch" in url_l:
        return "Facebook"
    if "youtube.com" in url_l or "youtu.be" in url_l:
        return "YouTube"
    return "Social"


def _record_status(record: ViewRecord | None) -> str:
    if record is None:
        return "Pending Sync"
    if int(record.views or 0) <= 0:
        return "Needs Verification"
    return "Synced"


def _compute_growth_metrics(records: list[ViewRecord], days: int):
    if not records:
        return {
            "start_views": 0,
            "end_views": 0,
            "growth_views": 0,
            "growth_pct": 0.0,
            "avg_daily_growth": 0.0,
        }

    sorted_records = sorted(records, key=lambda r: r.date)
    end_record = sorted_records[-1]
    end_date = _safe_date(end_record.date)
    if end_date is None:
        return {
            "start_views": end_record.views,
            "end_views": end_record.views,
            "growth_views": 0,
            "growth_pct": 0.0,
            "avg_daily_growth": 0.0,
        }

    target_date = end_date - timedelta(days=days)
    baseline = None

    for r in sorted_records:
        rd = _safe_date(r.date)
        if rd is None:
            continue
        if rd <= target_date:
            baseline = r
        else:
            break

    if baseline is None:
        baseline = sorted_records[0]

    start_views = baseline.views
    end_views = end_record.views
    growth_views = end_views - start_views
    growth_pct = round((growth_views / start_views) * 100, 2) if start_views > 0 else 0.0

    baseline_date = _safe_date(baseline.date)
    span_days = (end_date - baseline_date).days if baseline_date else days
    span_days = max(span_days, 1)
    avg_daily_growth = round(growth_views / span_days, 2)

    return {
        "start_views": start_views,
        "end_views": end_views,
        "growth_views": growth_views,
        "growth_pct": growth_pct,
        "avg_daily_growth": avg_daily_growth,
    }


def _build_video_summaries(videos: list[Video], all_records: list[ViewRecord]):
    records_by_video = defaultdict(list)
    for r in all_records:
        records_by_video[r.video_id].append(r)

    summaries = []
    for v in videos:
        v_records = records_by_video.get(v.id, [])
        if not v_records:
            summaries.append({
                "video_id": v.id,
                "title": v.title,
                "url": v.url,
                "thumbnail": v.thumbnail,
                "platform": _platform_label(v.url),
                "status": "Pending Sync",
                "current_views": None,
                "growth_7d": 0,
                "growth_7d_pct": 0.0,
                "growth_30d": 0,
                "growth_30d_pct": 0.0,
                "growth_365d": 0,
                "growth_365d_pct": 0.0,
                "avg_daily_7d": 0.0,
                "avg_daily_30d": 0.0,
                "avg_daily_365d": 0.0,
            })
            continue

        latest = max(v_records, key=lambda r: r.date)
        m7 = _compute_growth_metrics(v_records, 7)
        m30 = _compute_growth_metrics(v_records, 30)
        m365 = _compute_growth_metrics(v_records, 365)

        summaries.append({
            "video_id": v.id,
            "title": v.title,
            "url": v.url,
            "thumbnail": v.thumbnail,
            "platform": _platform_label(v.url),
            "status": _record_status(latest),
            "current_views": latest.views if latest.views > 0 else None,
            "growth_7d": m7["growth_views"],
            "growth_7d_pct": m7["growth_pct"],
            "growth_30d": m30["growth_views"],
            "growth_30d_pct": m30["growth_pct"],
            "growth_365d": m365["growth_views"],
            "growth_365d_pct": m365["growth_pct"],
            "avg_daily_7d": m7["avg_daily_growth"],
            "avg_daily_30d": m30["avg_daily_growth"],
            "avg_daily_365d": m365["avg_daily_growth"],
        })

    return summaries


def _build_current_rows(videos: list[Video], all_records: list[ViewRecord]):
    records_by_video = defaultdict(list)
    for r in all_records:
        records_by_video[r.video_id].append(r)

    rows = []
    for v in videos:
        records = records_by_video.get(v.id, [])
        latest = max(records, key=lambda r: r.date) if records else None
        rows.append({
            "date": latest.date if latest else "",
            "title": v.title,
            "url": v.url,
            "thumbnail": v.thumbnail,
            "platform": _platform_label(v.url),
            "views": latest.views if latest and latest.views > 0 else None,
            "growth": latest.growth if latest else 0,
            "status": _record_status(latest),
        })

    rows.sort(key=lambda x: (x["status"] == "Synced", x["date"]), reverse=True)
    return rows


def get_export_data(db: Session, current_user: User, days: int):
    videos = db.query(Video).filter(Video.user_id == current_user.id).all()
    video_ids = [v.id for v in videos]

    if not video_ids:
        return [], [], []

    cutoff_date = (datetime.utcnow().date() - timedelta(days=days)).isoformat()
    all_records = db.query(ViewRecord).filter(
        ViewRecord.video_id.in_(video_ids), ViewRecord.date >= cutoff_date
    ).order_by(ViewRecord.date.desc()).all()
    full_records = db.query(ViewRecord).filter(
        ViewRecord.video_id.in_(video_ids)
    ).all()

    video_map = {v.id: v for v in videos}
    full_by_video = defaultdict(list)
    for r in full_records:
        full_by_video[r.video_id].append(r)

    rows = []
    for r in all_records:
        v = video_map.get(r.video_id)
        if v:
            rows.append({
                "date": r.date,
                "title": v.title,
                "url": v.url,
                "views": r.views,
                "growth": r.growth,
                "status": "Synced",
            })

    period_summaries = []
    for v in videos:
        metrics = _compute_growth_metrics(full_by_video.get(v.id, []), days)
        period_summaries.append({
            "title": v.title,
            "url": v.url,
            "current_views": metrics["end_views"],
            "period_start_views": metrics["start_views"],
            "period_growth_views": metrics["growth_views"],
            "period_growth_pct": metrics["growth_pct"],
            "avg_daily_growth": metrics["avg_daily_growth"],
        })

    return videos, rows, period_summaries


@router.get("/overview")
def get_analytics_overview(
    days: int = 30,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if current_user.is_super_admin:
        raise HTTPException(status_code=403, detail="Use super admin reports endpoint")

    videos = db.query(Video).filter(Video.user_id == current_user.id).all()
    video_ids = [v.id for v in videos]

    if not video_ids:
        return {
            "total_videos": 0,
            "total_views": 0,
            "views_today": 0,
            "highest_video": None,
            "trends": [],
            "table_data": [],
            "video_summaries": [],
        }

    cutoff_date = (datetime.utcnow().date() - timedelta(days=days)).isoformat()
    today_str = datetime.utcnow().date().isoformat()

    all_records = db.query(ViewRecord).filter(
        ViewRecord.video_id.in_(video_ids), ViewRecord.date >= cutoff_date
    ).all()
    full_records = db.query(ViewRecord).filter(
        ViewRecord.video_id.in_(video_ids)
    ).all()

    if not all_records:
        return {
            "total_videos": len(videos),
            "total_views": 0,
            "views_today": 0,
            "highest_video": None,
            "trends": [],
            "table_data": _build_current_rows(videos, []),
            "video_summaries": _build_video_summaries(videos, []),
        }

    latest_records = {}
    records_for_latest = full_records if full_records else all_records
    for v in videos:
        records = [r for r in records_for_latest if r.video_id == v.id]
        if records:
            latest = max(records, key=lambda r: r.date)
            latest_records[v.id] = latest

    total_views = sum(r.views for r in latest_records.values())

    today_records = [r for r in all_records if r.date == today_str]
    previous_records = db.query(ViewRecord).filter(
        ViewRecord.video_id.in_(video_ids),
        ViewRecord.date < today_str
    ).all()

    latest_previous_by_video = {}
    for r in previous_records:
        existing = latest_previous_by_video.get(r.video_id)
        if existing is None or r.date > existing.date:
            latest_previous_by_video[r.video_id] = r

    views_today = 0
    for r in today_records:
        prev = latest_previous_by_video.get(r.video_id)
        prev_views = prev.views if prev else 0
        views_today += (r.views - prev_views)

    highest = max(latest_records.items(), key=lambda x: x[1].views) if latest_records else None
    highest_video = None
    if highest:
        vid = next(v for v in videos if v.id == highest[0])
        highest_video = {"title": vid.title, "views": highest[1].views}

    date_views = {}
    for r in all_records:
        if r.date not in date_views:
            date_views[r.date] = 0
        date_views[r.date] += r.views

    trends = sorted(
        [{"date": d, "views": v} for d, v in date_views.items()],
        key=lambda x: x["date"],
    )

    table_data = _build_current_rows(videos, full_records if full_records else all_records)

    return {
        "total_videos": len(videos),
        "total_views": total_views,
        "views_today": views_today,
        "highest_video": highest_video,
        "trends": trends,
        "table_data": table_data,
        "video_summaries": _build_video_summaries(videos, full_records if full_records else all_records),
    }


@router.get("/super-admin/overview")
def super_admin_overview(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if not current_user.is_super_admin:
        raise HTTPException(status_code=403, detail="Super admin access required")

    total_companies = db.query(User).filter(User.is_super_admin == False).count()
    active_companies = db.query(User).filter(
        User.is_super_admin == False, User.status == "Active"
    ).count()
    total_videos = db.query(Video).count()
    total_records = db.query(ViewRecord).count()

    total_views = 0
    if total_records > 0:
        all_records = db.query(ViewRecord).all()
        video_max = {}
        for r in all_records:
            if r.video_id not in video_max or r.views > video_max[r.video_id]:
                video_max[r.video_id] = r.views
        total_views = sum(video_max.values())

    return {
        "total_companies": total_companies,
        "active_companies": active_companies,
        "total_videos": total_videos,
        "total_views": total_views,
    }


@router.get("/super-admin/reports")
def super_admin_reports(
    days: int = 30,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if not current_user.is_super_admin:
        raise HTTPException(status_code=403, detail="Super admin access required")

    companies = db.query(User).filter(User.is_super_admin == False).all()
    company_ids = [c.id for c in companies]
    if not company_ids:
        return {"company_summaries": [], "detail_rows": []}

    videos = db.query(Video).filter(Video.user_id.in_(company_ids)).all()
    videos_by_company = defaultdict(list)
    video_map = {}
    for v in videos:
        videos_by_company[v.user_id].append(v)
        video_map[v.id] = v

    video_ids = [v.id for v in videos]
    if not video_ids:
        company_summaries = [{
            "company_id": c.id,
            "company_name": c.company_name,
            "company_email": c.email,
            "total_videos": 0,
            "total_views": 0,
            "views_today": 0,
            "highest_video": None,
        } for c in companies]
        return {"company_summaries": company_summaries, "detail_rows": []}

    cutoff_date = (datetime.utcnow().date() - timedelta(days=days)).isoformat()
    today_str = datetime.utcnow().date().isoformat()

    all_records = db.query(ViewRecord).filter(ViewRecord.video_id.in_(video_ids)).all()
    filtered_records = [r for r in all_records if r.date >= cutoff_date]

    records_by_video = defaultdict(list)
    filtered_by_video = defaultdict(list)
    for r in all_records:
        records_by_video[r.video_id].append(r)
    for r in filtered_records:
        filtered_by_video[r.video_id].append(r)

    company_summaries = []
    detail_rows = []

    for company in companies:
        company_videos = videos_by_company.get(company.id, [])
        latest_views_total = 0
        views_today_total = 0
        highest_video = None
        highest_views = -1

        for v in company_videos:
            full_v_records = records_by_video.get(v.id, [])
            period_records = filtered_by_video.get(v.id, [])

            if full_v_records:
                latest = max(full_v_records, key=lambda r: r.date)
                latest_views_total += latest.views
                if latest.views > highest_views:
                    highest_views = latest.views
                    highest_video = {"title": v.title, "views": latest.views}

            today_record = next((r for r in period_records if r.date == today_str), None)
            if today_record:
                prev = None
                for r in full_v_records:
                    if r.date < today_str and (prev is None or r.date > prev.date):
                        prev = r
                prev_views = prev.views if prev else 0
                views_today_total += (today_record.views - prev_views)

            for r in period_records:
                detail_rows.append({
                    "company_id": company.id,
                    "company_name": company.company_name,
                    "date": r.date,
                    "title": v.title,
                    "url": v.url,
                    "thumbnail": v.thumbnail,
                    "views": r.views,
                    "growth": r.growth,
                    "status": "Synced",
                })

        company_summaries.append({
            "company_id": company.id,
            "company_name": company.company_name,
            "company_email": company.email,
            "total_videos": len(company_videos),
            "total_views": latest_views_total,
            "views_today": views_today_total,
            "highest_video": highest_video,
        })

    detail_rows.sort(key=lambda x: x["date"], reverse=True)
    company_summaries.sort(key=lambda x: x["total_views"], reverse=True)

    return {
        "company_summaries": company_summaries,
        "detail_rows": detail_rows[:1200],
    }


def format_date(date_str: str) -> str:
    try:
        return datetime.strptime(date_str, "%Y-%m-%d").strftime("%b %d, %Y")
    except Exception:
        return date_str


@router.get("/export/csv")
def export_csv(
    days: int = 30,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if current_user.is_super_admin:
        raise HTTPException(status_code=403, detail="Super admins cannot export CSV")

    videos, rows, summaries = get_export_data(db, current_user, days)
    if not rows:
        raise HTTPException(status_code=404, detail="No data to export")

    output = io.StringIO()
    writer = csv.writer(output)

    writer.writerow(["ViewPulse Analytics Report"])
    writer.writerow([f"Company: {current_user.company_name}"])
    writer.writerow([f"Generated: {datetime.utcnow().strftime('%B %d, %Y %I:%M %p')}"])
    writer.writerow([f"Period: Last {days} days"])
    writer.writerow([])

    writer.writerow(["Date", "Video Title", "YouTube URL", "Views", "Daily Growth %", "Status"])

    for r in rows:
        writer.writerow([
            format_date(r["date"]),
            r["title"],
            r["url"],
            r["views"],
            f"+{r['growth']}%",
            r["status"],
        ])

    writer.writerow([])
    writer.writerow([f"Period Summary ({days} Days)"])
    writer.writerow(["Video Title", "Start Views", "Current Views", "Growth", "Growth %", "Avg Daily Growth"])
    for s in summaries:
        writer.writerow([
            s["title"],
            s["period_start_views"],
            s["current_views"],
            s["period_growth_views"],
            f"{s['period_growth_pct']}%",
            s["avg_daily_growth"],
        ])

    output.seek(0)
    filename = f"viewpulse-report-{datetime.utcnow().strftime('%Y-%m-%d')}.csv"
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/export/excel")
def export_excel(
    days: int = 30,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if current_user.is_super_admin:
        raise HTTPException(status_code=403, detail="Super admins cannot export Excel")

    videos, rows, summaries = get_export_data(db, current_user, days)
    if not rows:
        raise HTTPException(status_code=404, detail="No data to export")

    wb = Workbook()

    ws = wb.active
    ws.title = "View Analytics"
    ws.sheet_properties.tabColor = "FF0033"

    header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="FF0033", end_color="FF0033", fill_type="solid")
    subheader_font = Font(name="Calibri", size=11, bold=True)
    normal_font = Font(name="Calibri", size=10)
    date_font = Font(name="Calibri", size=10)
    views_font = Font(name="Calibri", size=10, bold=True)
    growth_font = Font(name="Calibri", size=10, color="00AA00")
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    ws.merge_cells("A1:F1")
    ws["A1"] = "ViewPulse Analytics Report"
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color="FF0033")
    ws["A1"].alignment = Alignment(horizontal="left")
    ws.row_dimensions[1].height = 35

    ws.merge_cells("A2:F2")
    ws["A2"] = f"Company: {current_user.company_name}    |    Generated: {datetime.utcnow().strftime('%B %d, %Y %I:%M %p')}    |    Period: Last {days} days"
    ws["A2"].font = Font(name="Calibri", size=10, color="666666")
    ws.row_dimensions[2].height = 22

    ws.merge_cells("A3:F3")
    ws["A3"] = f"Total Videos: {len(set(r['title'] for r in rows))}    |    Total Records: {len(rows)}"
    ws["A3"].font = Font(name="Calibri", size=10, bold=True, color="333333")
    ws.row_dimensions[3].height = 22

    for r_idx in range(1, 7):
        ws.column_dimensions[get_column_letter(r_idx)].width = 0

    headers = ["Date", "Video Title", "YouTube URL", "Total Views", "Daily Growth", "Status"]
    col_widths = [18, 50, 52, 16, 16, 14]
    start_row = 5

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[start_row].height = 28

    for row_idx, r in enumerate(rows, start_row + 1):
        date_cell = ws.cell(row=row_idx, column=1, value=format_date(r["date"]))
        date_cell.font = date_font
        date_cell.alignment = Alignment(horizontal="center")
        date_cell.border = thin_border

        title_cell = ws.cell(row=row_idx, column=2, value=r["title"])
        title_cell.font = normal_font
        title_cell.alignment = Alignment(vertical="center", wrap_text=True)
        title_cell.border = thin_border

        url_cell = ws.cell(row=row_idx, column=3, value=r["url"])
        url_cell.font = Font(name="Calibri", size=9, color="0066CC")
        url_cell.alignment = Alignment(vertical="center")
        url_cell.border = thin_border

        views_cell = ws.cell(row=row_idx, column=4, value=r["views"])
        views_cell.font = views_font
        views_cell.alignment = Alignment(horizontal="center")
        views_cell.number_format = "#,##0"
        views_cell.border = thin_border

        growth_cell = ws.cell(row=row_idx, column=5, value=f"+{r['growth']}%")
        growth_cell.font = growth_font
        growth_cell.alignment = Alignment(horizontal="center")
        growth_cell.border = thin_border

        status_cell = ws.cell(row=row_idx, column=6, value=r["status"])
        status_cell.font = Font(name="Calibri", size=10, color="009900")
        status_cell.alignment = Alignment(horizontal="center")
        status_cell.border = thin_border

        if row_idx % 2 == 0:
            for col_idx in range(1, 7):
                ws.cell(row=row_idx, column=col_idx).fill = PatternFill(
                    start_color="F8F9FA", end_color="F8F9FA", fill_type="solid"
                )

    ws.auto_filter.ref = f"A{start_row}:F{start_row + len(rows)}"
    ws.freeze_panes = f"A{start_row + 1}"

    ws_summary = wb.create_sheet("Summary")
    ws_summary.sheet_properties.tabColor = "10B981"

    ws_summary.merge_cells("A1:B1")
    ws_summary["A1"] = "Video Summary"
    ws_summary["A1"].font = Font(name="Calibri", size=16, bold=True, color="FF0033")
    ws_summary.row_dimensions[1].height = 35

    summary_headers = ["Video Title", "Start Views", "Current Views", "Growth", "Growth %", "Avg Daily"]
    summary_widths = [44, 14, 14, 14, 12, 12]

    for col_idx, (header, width) in enumerate(zip(summary_headers, summary_widths), 1):
        cell = ws_summary.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        ws_summary.column_dimensions[get_column_letter(col_idx)].width = width

    ws_summary.row_dimensions[3].height = 28

    for row_idx, s in enumerate(summaries, 4):
        ws_summary.cell(row=row_idx, column=1, value=s["title"]).font = normal_font
        ws_summary.cell(row=row_idx, column=1).border = thin_border
        ws_summary.cell(row=row_idx, column=1).alignment = Alignment(vertical="center", wrap_text=True)

        start_cell = ws_summary.cell(row=row_idx, column=2, value=s["period_start_views"])
        start_cell.font = views_font
        start_cell.alignment = Alignment(horizontal="center")
        start_cell.number_format = "#,##0"
        start_cell.border = thin_border

        current_cell = ws_summary.cell(row=row_idx, column=3, value=s["current_views"])
        current_cell.font = views_font
        current_cell.alignment = Alignment(horizontal="center")
        current_cell.number_format = "#,##0"
        current_cell.border = thin_border

        growth_cell = ws_summary.cell(row=row_idx, column=4, value=s["period_growth_views"])
        growth_cell.font = views_font
        growth_cell.alignment = Alignment(horizontal="center")
        growth_cell.number_format = "#,##0"
        growth_cell.border = thin_border

        pct_cell = ws_summary.cell(row=row_idx, column=5, value=f"{s['period_growth_pct']}%")
        pct_cell.font = growth_font
        pct_cell.alignment = Alignment(horizontal="center")
        pct_cell.border = thin_border

        avg_cell = ws_summary.cell(row=row_idx, column=6, value=s["avg_daily_growth"])
        avg_cell.font = normal_font
        avg_cell.alignment = Alignment(horizontal="center")
        avg_cell.border = thin_border

    ws_summary.auto_filter.ref = f"A3:F{3 + len(summaries)}"
    ws_summary.freeze_panes = "A4"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"viewpulse-report-{datetime.utcnow().strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/export/pdf")
def export_pdf(
    days: int = 30,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if current_user.is_super_admin:
        raise HTTPException(status_code=403, detail="Super admins cannot export PDF")

    videos, rows, summaries = get_export_data(db, current_user, days)
    if not rows:
        raise HTTPException(status_code=404, detail="No data to export")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=0.7 * inch,
        bottomMargin=0.5 * inch,
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
    )

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ReportTitle",
        fontName="Helvetica-Bold",
        fontSize=22,
        textColor=colors.HexColor("#FF0033"),
        spaceAfter=4,
    ))
    styles.add(ParagraphStyle(
        name="ReportSub",
        fontName="Helvetica",
        fontSize=9,
        textColor=colors.HexColor("#666666"),
        spaceAfter=2,
    ))
    styles.add(ParagraphStyle(
        name="TableHeader",
        fontName="Helvetica-Bold",
        fontSize=8,
        textColor=colors.white,
        alignment=1,
    ))
    styles.add(ParagraphStyle(
        name="TableCell",
        fontName="Helvetica",
        fontSize=7.5,
        leading=10,
    ))
    styles.add(ParagraphStyle(
        name="TableCellCenter",
        fontName="Helvetica",
        fontSize=7.5,
        alignment=1,
        leading=10,
    ))
    styles.add(ParagraphStyle(
        name="TableCellBold",
        fontName="Helvetica-Bold",
        fontSize=7.5,
        alignment=1,
        leading=10,
    ))
    styles.add(ParagraphStyle(
        name="TableCellGreen",
        fontName="Helvetica-Bold",
        fontSize=7.5,
        textColor=colors.HexColor("#00AA00"),
        alignment=1,
    ))
    styles.add(ParagraphStyle(
        name="TableCellBlue",
        fontName="Helvetica",
        fontSize=6.5,
        textColor=colors.HexColor("#0066CC"),
        leading=10,
    ))

    elements = []

    elements.append(Paragraph("ViewPulse Analytics Report", styles["ReportTitle"]))
    elements.append(Paragraph(
        f"Company: {current_user.company_name}  |  Generated: {datetime.utcnow().strftime('%B %d, %Y %I:%M %p')}  |  Period: Last {days} days",
        styles["ReportSub"],
    ))
    elements.append(Paragraph(
        f"Total Videos: {len(set(r['title'] for r in rows))}  |  Total Records: {len(rows)}",
        ParagraphStyle("StatsLine", fontName="Helvetica-Bold", fontSize=9, textColor=colors.HexColor("#333333"), spaceAfter=12),
    ))

    col_widths = [70, 150, 120, 55, 50, 45]
    headers = ["Date", "Video Title", "URL", "Views", "Growth", "Status"]

    table_data = []

    header_row = [Paragraph(h, styles["TableHeader"]) for h in headers]
    table_data.append(header_row)

    for r in rows:
        row_data = [
            Paragraph(format_date(r["date"]), styles["TableCellCenter"]),
            Paragraph(r["title"], styles["TableCell"]),
            Paragraph(r["url"], styles["TableCellBlue"]),
            Paragraph(f'{r["views"]:,}', styles["TableCellBold"]),
            Paragraph(f'+{r["growth"]}%', styles["TableCellGreen"]),
            Paragraph(r["status"], ParagraphStyle(
                "StatusCell", fontName="Helvetica-Bold", fontSize=7.5,
                textColor=colors.HexColor("#009900"), alignment=1,
            )),
        ]
        table_data.append(row_data)

    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FF0033")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F9FA")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E0E0E0")),
        ("TOPPADDING", (0, 1), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))

    elements.append(table)

    elements.append(Spacer(1, 20))

    elements.append(Paragraph("Video Summary", ParagraphStyle(
        "SummaryTitle", fontName="Helvetica-Bold", fontSize=14,
        textColor=colors.HexColor("#FF0033"), spaceAfter=8,
    )))

    summary_headers = ["Video Title", "Start", "Current", "Growth", "Growth %", "Avg/Day"]
    summary_widths = [215, 55, 55, 55, 50, 50]

    summary_data = []
    summary_data.append([Paragraph(h, styles["TableHeader"]) for h in summary_headers])

    for s in summaries:
        summary_data.append([
            Paragraph(s["title"], styles["TableCell"]),
            Paragraph(f'{s["period_start_views"]:,}', styles["TableCellCenter"]),
            Paragraph(f'{s["current_views"]:,}', styles["TableCellBold"]),
            Paragraph(f'{s["period_growth_views"]:,}', styles["TableCellBold"]),
            Paragraph(f'{s["period_growth_pct"]}%', styles["TableCellGreen"]),
            Paragraph(str(s["avg_daily_growth"]), styles["TableCellCenter"]),
        ])

    summary_table = Table(summary_data, colWidths=summary_widths, repeatRows=1)
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#10B981")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F9FA")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E0E0E0")),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))

    elements.append(summary_table)

    elements.append(Spacer(1, 15))
    elements.append(Paragraph(
        f"ViewPulse Analytics  |  Generated on {datetime.utcnow().strftime('%B %d, %Y')}",
        ParagraphStyle("Footer", fontName="Helvetica", fontSize=7, textColor=colors.HexColor("#999999"), alignment=1),
    ))

    doc.build(elements)
    buffer.seek(0)

    filename = f"viewpulse-report-{datetime.utcnow().strftime('%Y-%m-%d')}.pdf"
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
